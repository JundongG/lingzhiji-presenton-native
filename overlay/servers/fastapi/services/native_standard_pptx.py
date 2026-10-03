# SPDX-License-Identifier: Apache-2.0
# Independent native-export add-on; no licensed export-core implementation is used.
"""Independent, bounded native PPTX exporter for Presenton's Standard canvas.

This module consumes resolved ``slide.ui`` dictionaries and optional caller-supplied
PNG/JPEG bytes, not HTML/PDF/template source or any licensed export pipeline.
It performs no networking, filesystem asset reads/writes, database changes, or
mutations of its inputs. References only index the supplied in-memory mapping.
All output uses the existing python-pptx and Pillow dependencies.

Coordinates and font sizes are canvas pixels (1280x720 at 96 DPI). Supported:
explicitly sized text/text lists, tables, native bar/line/donut charts, container
rectangles, polygon/ellipse vectors (including uniform rounded rectangles),
positioned groups, supplied native PNG/JPEG pictures, and basic fixed flow/grid
layouts. Container-bounded child text/flow sizes can use their parent's bounds.
Background images, inline Markdown, LaTeX, infographics, unbounded dynamic sizing,
advanced or manual flow layouts, clipped descendants and unsupported transforms
fail with a path by default. ``unsupported='warn'`` explicitly permits
omission and returns structured warnings; it never rasterizes missing content.
Native Office charts are editable and include their own embedded XLSX workbook.
The safe literal workbook writer uses private extension points verified against
python-pptx 1.0.2; pin that version and rerun these tests before upgrading it.
Office's chart styling and text shaping can differ from the canvas renderer.
"""

from __future__ import annotations

from contextlib import contextmanager
from dataclasses import dataclass
from io import BytesIO
import math
import re
from typing import Any, Literal, Mapping, Sequence

from PIL import Image as PillowImage, UnidentifiedImageError

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.chart.xlsx import CategoryWorkbookWriter
from xlsxwriter import Workbook
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

CANVAS_WIDTH = 1280
CANVAS_HEIGHT = 720
EMU_PER_PIXEL = 9525
DEFAULT_FONT = {"family": "Arial", "size": 18, "color": "1A2B45"}
DEFAULT_CHART_COLORS = ("7F22FE",)
MAX_SLIDES = 200
MAX_ELEMENTS = 10000
MAX_DEPTH = 32
MAX_INPUT_NODES = 100000
MAX_TEXT_CHARACTERS = 2000000
MAX_TABLE_CELLS = 10000
MAX_CHART_VALUES = 100000
MAX_VECTOR_POINTS = 20000
MAX_WORKBOOK_STRING_LENGTH = 32767
MAX_IMAGE_ASSETS = 100
MAX_IMAGE_BYTES = 8 * 1024 * 1024
MAX_TOTAL_IMAGE_BYTES = 32 * 1024 * 1024
MAX_IMAGE_PIXELS = 16000000
MAX_TOTAL_IMAGE_PIXELS = 64000000
MAX_IMAGE_SIDE = 8192


@dataclass(frozen=True)
class ExportWarning:
    code: str
    path: str
    message: str


@dataclass(frozen=True)
class StandardPptxExport:
    pptx_bytes: bytes
    warnings: tuple[ExportWarning, ...]


class NativeStandardExportError(ValueError):
    """An invalid or unsupported input; safe to display as a validation error."""

    def __init__(self, message: str, *, path: str, code: str = "invalid_input"):
        self.path = path
        self.code = code
        super().__init__(f"{path}: {message}")


class _SkipElement(Exception):
    pass


def _record(value: Any, path: str) -> Mapping[str, Any]:
    if not isinstance(value, Mapping):
        raise NativeStandardExportError("Expected an object", path=path)
    return value


def _font_values(value: Any, path: str) -> dict[str, Any]:
    """Optional schema properties may be explicit null; they inherit defaults."""
    if value is None:
        return {}
    return {key: item for key, item in _record(value, path).items() if item is not None}


def _number(value: Any, path: str, *, minimum: float | None = None) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise NativeStandardExportError("Expected a finite number", path=path)
    # Check integer magnitude before math.isfinite coerces enormous integers
    # to float (which can otherwise raise an unstructured OverflowError).
    if abs(value) > 1e9:
        raise NativeStandardExportError("Number exceeds export limit", path=path)
    if not math.isfinite(value):
        raise NativeStandardExportError("Expected a finite number", path=path)
    if minimum is not None and value < minimum:
        raise NativeStandardExportError(f"Expected a number >= {minimum}", path=path)
    return float(value)


def _px(value: float) -> int:
    return round(value * EMU_PER_PIXEL)


def _color(value: Any, path: str) -> RGBColor:
    if not isinstance(value, str):
        raise NativeStandardExportError("Expected a hex color", path=path)
    value = value.removeprefix("#")
    if re.fullmatch(r"[0-9a-fA-F]{3}", value):
        value = "".join(char * 2 for char in value)
    if not re.fullmatch(r"[0-9a-fA-F]{6}", value):
        raise NativeStandardExportError("Only RGB hex colors are supported", path=path)
    return RGBColor.from_string(value.upper())


def _opacity(solid_fill: Any, opacity: Any, path: str) -> None:
    value = _number(opacity, path, minimum=0)
    if value > 1:
        raise NativeStandardExportError("Opacity must be between 0 and 1", path=path)
    rgb = solid_fill.find("{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
    if rgb is not None:
        alpha = OxmlElement("a:alpha")
        alpha.set("val", str(round(value * 100000)))
        rgb.append(alpha)


def _point(element: Mapping[str, Any], path: str) -> tuple[float, float]:
    position = _record(element.get("position") or {}, path + ".position")
    return (_number(position.get("x", 0), path + ".position.x"),
            _number(position.get("y", 0), path + ".position.y"))


def _size(element: Mapping[str, Any], path: str) -> tuple[float, float]:
    size = _record(element.get("size"), path + ".size")
    return (_number(size.get("width"), path + ".size.width", minimum=0.01),
            _number(size.get("height"), path + ".size.height", minimum=0.01))


def _padding(value: Any, path: str) -> tuple[float, float, float, float]:
    if value is None:
        return 0, 0, 0, 0
    if isinstance(value, (float, int)):
        size = _number(value, path, minimum=0)
        return size, size, size, size
    value = _record(value, path)
    return tuple(_number(value.get(key, 0), path + "." + key, minimum=0)
                 for key in ("top", "right", "bottom", "left"))


class _LiteralWorkbookWriter(CategoryWorkbookWriter):
    """Keep category/series strings literal, never formulas or hyperlinks.

    This exporter-local extension uses python-pptx's existing category layout
    and reference generation. It does not monkeypatch global library behavior.
    The embedded-workbook regression tests cover this private extension point.
    """

    @contextmanager
    def _open_worksheet(self, xlsx_file: BytesIO):
        with Workbook(xlsx_file, {"in_memory": True, "strings_to_formulas": False,
                                  "strings_to_urls": False, "strings_to_numbers": False}) as workbook:
            yield workbook, workbook.add_worksheet()


class _LiteralChartData(CategoryChartData):
    @property
    def _workbook_writer(self):
        return _LiteralWorkbookWriter(self)


def _has_inline_markup(text: str) -> bool:
    # The canvas interprets these delimiters even inside ordinary runs. This
    # bounded exporter rejects marked-up runs rather than emitting raw markup
    # and claiming that it preserved their rendered appearance.
    return re.search(r"(\*\*|__|[*_]).+?\1", text, re.DOTALL) is not None


def _validate_complexity(slide_uis: Sequence[Mapping[str, Any]]) -> None:
    """Bound aggregate work before creating shapes or embedded workbooks."""
    totals = {"nodes": 0, "characters": 0, "cells": 0, "values": 0, "points": 0}
    limits = {"nodes": MAX_INPUT_NODES, "characters": MAX_TEXT_CHARACTERS,
              "cells": MAX_TABLE_CELLS, "values": MAX_CHART_VALUES, "points": MAX_VECTOR_POINTS}
    ancestors: set[int] = set()

    def add(kind: str, amount: int, path: str) -> None:
        totals[kind] += amount
        if totals[kind] > limits[kind]:
            raise NativeStandardExportError(f"Aggregate {kind} limit exceeded ({limits[kind]})", path=path, code="export_limit")

    def visit(value: Any, path: str, depth: int) -> None:
        add("nodes", 1, path)
        if depth > MAX_DEPTH * 4:
            raise NativeStandardExportError("Input nesting limit exceeded", path=path, code="export_limit")
        if isinstance(value, str):
            add("characters", len(value), path)
        elif isinstance(value, (Mapping, list, tuple)):
            if id(value) in ancestors:
                raise NativeStandardExportError("Cyclic input is not supported", path=path)
            ancestors.add(id(value))
            if isinstance(value, Mapping):
                if value.get("type") == "table":
                    columns, rows = value.get("columns"), value.get("rows")
                    cells = len(columns) if isinstance(columns, list) else 0
                    if isinstance(rows, list):
                        cells += sum(len(row) for row in rows if isinstance(row, list))
                    add("cells", cells, path)
                elif value.get("type") == "chart":
                    series = value.get("series")
                    if isinstance(series, list):
                        add("values", sum(len(item["values"]) for item in series if isinstance(item, Mapping) and isinstance(item.get("values"), list)), path)
                elif value.get("type") == "vector":
                    points = value.get("points")
                    add("points", len(points) if isinstance(points, list) else 0, path)
                for key, item in value.items():
                    if not isinstance(key, str):
                        raise NativeStandardExportError("Object keys must be strings", path=path)
                    visit(item, path + "." + key, depth + 1)
            else:
                for index, item in enumerate(value):
                    visit(item, f"{path}[{index}]", depth + 1)
            ancestors.remove(id(value))

    for index, ui in enumerate(slide_uis):
        visit(ui, f"slides[{index}].ui", 0)


@dataclass(frozen=True)
class _ImageAsset:
    data: bytes
    width: int
    height: int


def _validate_image_assets(assets: Mapping[str, bytes] | None) -> dict[str, _ImageAsset]:
    """Validate supplied bytes only. Never interpret references as paths/URLs."""
    if assets is None:
        return {}
    assets = _record(assets, "image_assets")
    if len(assets) > MAX_IMAGE_ASSETS:
        raise NativeStandardExportError("Image asset count limit exceeded", path="image_assets", code="export_limit")
    total = 0
    decoded_pixels = 0
    validated = {}
    for reference, blob in assets.items():
        path = f"image_assets[{reference!r}]"
        if not isinstance(reference, str) or not reference or len(reference) > 4096:
            raise NativeStandardExportError("Image references must be non-empty strings up to 4096 characters", path="image_assets", code="invalid_image_asset")
        if not isinstance(blob, bytes):
            raise NativeStandardExportError("Image assets must be supplied as immutable bytes, never paths or URLs", path=path, code="invalid_image_asset")
        total += len(blob)
        if len(blob) > MAX_IMAGE_BYTES or total > MAX_TOTAL_IMAGE_BYTES:
            raise NativeStandardExportError("Image byte limit exceeded", path=path, code="export_limit")
        try:
            with PillowImage.open(BytesIO(blob), formats=("PNG", "JPEG")) as decoded:
                width, height = decoded.size
                decoded_pixels += width * height
                if max(width, height) > MAX_IMAGE_SIDE or width * height > MAX_IMAGE_PIXELS or decoded_pixels > MAX_TOTAL_IMAGE_PIXELS:
                    raise NativeStandardExportError("Decoded image dimensions or aggregate pixel budget exceed limits", path=path, code="export_limit")
                if getattr(decoded, "n_frames", 1) != 1:
                    raise NativeStandardExportError("Animated images are not supported", path=path, code="invalid_image_asset")
                if decoded.getexif().get(274, 1) != 1:
                    raise NativeStandardExportError("EXIF-rotated images require caller-normalized pixels", path=path, code="invalid_image_asset")
                decoded.load()
        except NativeStandardExportError:
            raise
        except (UnidentifiedImageError, OSError, ValueError, PillowImage.DecompressionBombError) as exc:
            raise NativeStandardExportError("Expected a valid non-animated PNG or JPEG image", path=path, code="invalid_image_asset") from exc
        validated[reference] = _ImageAsset(blob, width, height)
    return validated


class _Exporter:
    def __init__(self, title: str, theme: Mapping[str, Any] | None, unsupported: str,
                 image_assets: Mapping[str, _ImageAsset]):
        self.presentation = Presentation()
        self.presentation.slide_width = _px(CANVAS_WIDTH)
        self.presentation.slide_height = _px(CANVAS_HEIGHT)
        self.presentation.core_properties.title = title
        self.presentation.core_properties.subject = "Native Standard canvas export"
        self.presentation.core_properties.author = ""
        self.presentation.core_properties.last_modified_by = ""
        self.presentation.core_properties.comments = ""
        self.theme = theme or {}
        self.image_assets = image_assets
        self.unsupported = unsupported
        self.warnings: list[ExportWarning] = []
        self.element_count = 0

    def reject(self, path: str, message: str, code: str = "unsupported_element") -> None:
        if self.unsupported == "error":
            raise NativeStandardExportError(message, path=path, code=code)
        self.warnings.append(ExportWarning(code, path, message + "; omitted"))
        raise _SkipElement()

    def appearance_warning(self, path: str, message: str) -> None:
        # These are visible limitations, not omitted content. The caller must
        # expose all warnings when integrating this result into an HTTP API.
        self.warnings.append(ExportWarning("appearance_difference", path, message))

    def run(self, slide_uis: Sequence[Mapping[str, Any]]) -> StandardPptxExport:
        for index, ui in enumerate(slide_uis):
            path = f"slides[{index}].ui"
            ui = _record(ui, path)
            slide = self.presentation.slides.add_slide(self.presentation.slide_layouts[6])
            for background_key in ("background_image", "backgroundImage"):
                if ui.get(background_key) is not None:
                    try:
                        self.reject(path + "." + background_key, "Background images are unsupported; no asset is fetched")
                    except _SkipElement:
                        pass
            colors = self.theme.get("data", self.theme).get("colors", {})
            background = ui.get("background") or colors.get("background") or "FFFFFF"
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = _color(background, path + ".background")
            elements = ui.get("elements", [])
            components = ui.get("components", [])
            if not isinstance(elements, list) or not isinstance(components, list):
                raise NativeStandardExportError("elements/components must be arrays", path=path)
            for component_index, component in enumerate(components):
                cpath = f"{path}.components[{component_index}]"
                component = _record(component, cpath)
                offset = _point(component, cpath)
                children = component.get("elements", [])
                if not isinstance(children, list):
                    raise NativeStandardExportError("elements must be an array", path=cpath)
                if component.get("opacity") is not None and _number(component["opacity"], cpath + ".opacity") != 1:
                    try:
                        self.reject(cpath + ".opacity", "Component-level opacity is unsupported")
                    except _SkipElement:
                        continue
                if any(component.get(key) for key in ("rotation", "flip_h", "flip_v")):
                    try:
                        self.reject(cpath, "Transformed components are not supported")
                    except _SkipElement:
                        continue
                for element_index, element in enumerate(children):
                    self.element(slide, element, offset, f"{cpath}.elements[{element_index}]")
            for element_index, element in enumerate(elements):
                self.element(slide, element, (0, 0), f"{path}.elements[{element_index}]")
        output = BytesIO()
        self.presentation.save(output)
        return StandardPptxExport(output.getvalue(), tuple(self.warnings))

    def element(self, slide: Any, element: Any, offset: tuple[float, float],
                path: str, depth: int = 0, box: tuple[float, float, float, float] | None = None) -> None:
        self.element_count += 1
        if depth > MAX_DEPTH or self.element_count > MAX_ELEMENTS:
            raise NativeStandardExportError("Export complexity limit exceeded", path=path)
        element = _record(element, path)
        # Roll back a partially emitted group on warn-mode failure rather than
        # leaving a misleading fragment of the unsupported element.
        before = set(slide.shapes._spTree)
        try:
            self._element(slide, element, offset, path, depth, box)
        except _SkipElement:
            for child in list(slide.shapes._spTree):
                if child not in before:
                    slide.shapes._spTree.remove(child)

    def _element(self, slide: Any, element: Mapping[str, Any], offset: tuple[float, float],
                 path: str, depth: int, box: tuple[float, float, float, float] | None) -> None:
        kind = element.get("type")
        if not isinstance(kind, str) or kind not in {"text", "text-list", "table", "chart", "container", "group", "vector", "flex", "grid", "image"}:
            self.reject(path, f"Element type {kind!r} is not supported")
        if kind != "image" and element.get("opacity") is not None and _number(element["opacity"], path + ".opacity") != 1:
            self.reject(path + ".opacity", "Element-level opacity is unsupported; font/fill/stroke opacity is supported")
        if element.get("shadow"):
            shadow = _record(element["shadow"], path + ".shadow")
            opacity = _number(shadow.get("opacity") if shadow.get("opacity") is not None else .2, path + ".shadow.opacity", minimum=0)
            if opacity > 1:
                raise NativeStandardExportError("Opacity must be between 0 and 1", path=path + ".shadow.opacity")
            blur = _number(shadow.get("blur") or 0, path + ".shadow.blur", minimum=0)
            dx = _number(shadow.get("offset_x") or 0, path + ".shadow.offset_x")
            dy = _number(shadow.get("offset_y") or 0, path + ".shadow.offset_y")
            if opacity > 0 and (blur or dx or dy):
                self.appearance_warning(path + ".shadow", "Shadows are not exported")
        if kind in {"group", "container", "flex", "grid"}:
            if any(element.get(key) for key in ("rotation", "flip_h", "flip_v")):
                self.reject(path, "Group/container rotation or mirroring is not supported")
        elif kind in {"table", "chart"} and any(element.get(key) for key in ("rotation", "flip_h", "flip_v")):
            self.reject(path, "Chart/table rotation or mirroring is not supported")
        if kind == "vector":
            self.vector(slide, element, offset, path)
            return
        x, y = _point(element, path)
        x, y = x + offset[0], y + offset[1]
        if kind == "group":
            if box is not None:
                x, y = box[:2]
            children = element.get("children", element.get("elements", []))
            self.children(slide, children, (x, y), path + ".children", depth)
            return
        if box is None:
            width, height = _size(element, path)
        else:
            x, y, width, height = box
        if kind in {"container", "flex", "grid"}:
            self.container(slide, element, (x, y, width, height), path, depth)
            return
        if x < 0 or y < 0 or x + width > CANVAS_WIDTH + 0.1 or y + height > CANVAS_HEIGHT + 0.1:
            self.appearance_warning(path, "Element extends beyond the slide bounds; its original geometry is preserved")
        if kind in {"text", "text-list"}:
            shape = slide.shapes.add_textbox(_px(x), _px(y), _px(width), _px(height))
            self.shape_style(shape, element, path)
            self.text(shape.text_frame, element, path, is_list=kind == "text-list")
        elif kind == "image":
            shape = self.picture(slide, element, (x, y, width, height), path)
        elif kind == "table":
            shape = self.table(slide, element, (x, y, width, height), path)
        else:
            shape = self.chart(slide, element, (x, y, width, height), path)
        shape.name = str(element.get("name") or kind)
        if element.get("rotation"):
            shape.rotation = _number(element["rotation"], path + ".rotation")
        if element.get("flip_h") or element.get("flip_v"):
            xfrm = shape._element.spPr.get_or_add_xfrm()
            if element.get("flip_h"):
                xfrm.set("flipH", "1")
            if element.get("flip_v"):
                xfrm.set("flipV", "1")

    def picture(self, slide: Any, element: Mapping[str, Any], box: tuple[float, float, float, float], path: str) -> Any:
        reference = element.get("data")
        if not isinstance(reference, str) or not reference:
            raise NativeStandardExportError("Image data must identify a supplied asset", path=path + ".data", code="invalid_image_asset")
        asset = self.image_assets.get(reference)
        if asset is None:
            self.reject(path + ".data", "Image reference has no supplied PNG/JPEG bytes", code="missing_image_asset")
        for clip_key in ("clip_path", "clippath", "clipPath"):
            clip = element.get(clip_key)
            if clip is not None and (not isinstance(clip, str) or clip.strip().lower() not in ("", "none")):
                self.reject(path + "." + clip_key, "Arbitrary image clip paths are unsupported")
        if element.get("color"):
            self.reject(path + ".color", "Icon recoloring is unsupported")
        crop_scale = _number(element.get("crop_scale") if element.get("crop_scale") is not None else 1, path + ".crop_scale", minimum=.1)
        if crop_scale != 1:
            self.reject(path + ".crop_scale", "Image zoom requires caller-resolved cropping")
        x, y, width, height = box
        fit = element.get("fit") or "contain"
        if not isinstance(fit, str) or fit not in {"fill", "cover", "contain"}:
            raise NativeStandardExportError("Unsupported image fit", path=path + ".fit")
        focus = []
        for axis in ("x", "y"):
            value = _number(element.get("focus_" + axis) if element.get("focus_" + axis) is not None else 50, path + ".focus_" + axis, minimum=0)
            if value > 100:
                raise NativeStandardExportError("Image focus must be between 0 and 100", path=path + ".focus_" + axis)
            focus.append(value / 100)
        draw_width, draw_height = width, height
        if fit == "contain":
            scale = min(width / asset.width, height / asset.height)
            draw_width, draw_height = asset.width * scale, asset.height * scale
            # Canvas image flips happen inside the original frame, so a
            # letterboxed picture's position must be mirrored as well as its
            # pixels. Native picture flips alone act around the smaller box.
            placed_focus_x = 1 - focus[0] if element.get("flip_h") else focus[0]
            placed_focus_y = 1 - focus[1] if element.get("flip_v") else focus[1]
            x, y = x + (width - draw_width) * placed_focus_x, y + (height - draw_height) * placed_focus_y
            if element.get("rotation") and focus != [.5, .5]:
                self.reject(path, "Rotated off-center contain images require resolved geometry")
        shape = slide.shapes.add_picture(BytesIO(asset.data), _px(x), _px(y), _px(draw_width), _px(draw_height))
        if fit == "cover":
            scale = max(width / asset.width, height / asset.height)
            crop_x = max(0, 1 - width / scale / asset.width)
            crop_y = max(0, 1 - height / scale / asset.height)
            shape.crop_left, shape.crop_right = crop_x * focus[0], crop_x * (1 - focus[0])
            shape.crop_top, shape.crop_bottom = crop_y * focus[1], crop_y * (1 - focus[1])
        radius = element.get("border_radius") or 0
        if isinstance(radius, Mapping):
            radii = [_number(radius.get(key, 0), path + ".border_radius." + key, minimum=0) for key in ("tl", "tr", "bl", "br")]
            if len(set(radii)) != 1:
                self.reject(path + ".border_radius", "Asymmetric picture corners are unsupported")
            radius = radii[0]
        radius = _number(radius, path + ".border_radius", minimum=0)
        if radius:
            if fit == "contain" and (abs(draw_width - width) > .01 or abs(draw_height - height) > .01):
                self.reject(path + ".border_radius", "Rounded contain-image frames require resolved visible geometry")
            geometry = shape._element.spPr.prstGeom
            geometry.set("prst", "roundRect")
            for old in list(geometry.avLst):
                geometry.avLst.remove(old)
            adjustment = OxmlElement("a:gd")
            adjustment.set("name", "adj")
            adjustment.set("fmla", f"val {round(min(.5, radius / min(width, height)) * 100000)}")
            geometry.avLst.append(adjustment)
        if element.get("opacity") is not None:
            opacity = _number(element["opacity"], path + ".opacity", minimum=0)
            if opacity > 1:
                raise NativeStandardExportError("Opacity must be between 0 and 1", path=path + ".opacity")
            alpha = OxmlElement("a:alphaModFix")
            alpha.set("amt", str(round(opacity * 100000)))
            shape._element.blipFill.blip.append(alpha)
        return shape

    def bounded_child_size(self, element: Mapping[str, Any], path: str,
                           fallback: tuple[float, float]) -> tuple[float, float]:
        """Resolve only parent-bounded optional sizes, never unbounded leaf text."""
        if element.get("size") is not None:
            return _size(element, path)
        kind = element.get("type")
        if kind in ("flex", "grid", "group", "image", "table", "chart"):
            return fallback
        if kind == "text":
            font = {**DEFAULT_FONT, **_font_values(element.get("font"), path + ".font")}
            size = _number(font.get("size", 18), path + ".font.size", minimum=.1)
            line_height = _number(font.get("line_height", 1.15), path + ".font.line_height", minimum=.1)
            if line_height > 2:
                self.reject(path + ".font.line_height", "Ambiguous line height in bounded text")
            runs = element.get("runs", [])
            if not isinstance(runs, list) or any(not isinstance(run, Mapping) or not isinstance(run.get("text"), str) for run in runs):
                raise NativeStandardExportError("Bounded text requires plain text runs", path=path + ".runs")
            content = "".join(run["text"] for run in runs)
            chars_per_line = max(1, math.floor(fallback[0] / max(1, size * .5)))
            lines = sum(max(1, math.ceil(len(line) / chars_per_line)) for line in re.split(r"\r?\n", content))
            return fallback[0], max(1, lines * size * line_height)
        return _size(element, path)

    def children(self, slide: Any, children: Any, offset: tuple[float, float], path: str, depth: int) -> None:
        if not isinstance(children, list):
            raise NativeStandardExportError("children must be an array", path=path)
        for index, child in enumerate(children):
            self.element(slide, child, offset, f"{path}[{index}]", depth + 1)

    def shape_style(self, shape: Any, element: Mapping[str, Any], path: str) -> None:
        # add_shape carries an Office theme style including a default shadow.
        # Explicitly suppress those effects: a source without a shadow should
        # never acquire one solely because of the blank PPTX template's theme.
        for effect_ref in shape._element.xpath("./p:style/a:effectRef"):
            effect_ref.set("idx", "0")
        for effects in shape._element.spPr.xpath("./a:effectLst"):
            shape._element.spPr.remove(effects)
        shape._element.spPr.append(OxmlElement("a:effectLst"))
        fill = element.get("fill")
        if fill:
            fill = _record(fill, path + ".fill")
            shape.fill.solid()
            shape.fill.fore_color.rgb = _color(fill.get("color"), path + ".fill.color")
            if fill.get("opacity") is not None:
                _opacity(shape.fill._xPr.solidFill, fill["opacity"], path + ".fill.opacity")
        else:
            shape.fill.background()
        stroke = element.get("stroke")
        if stroke is not None:
            stroke = _record(stroke, path + ".stroke")
        if stroke and _number(stroke.get("width", 0), path + ".stroke.width", minimum=0) > 0:
            shape.line.color.rgb = _color(stroke.get("color"), path + ".stroke.color")
            shape.line.width = _px(_number(stroke["width"], path + ".stroke.width", minimum=0))
            if stroke.get("opacity") is not None:
                _opacity(shape.line._get_or_add_ln().solidFill, stroke["opacity"], path + ".stroke.opacity")
            for key in ("dash", "start_marker", "end_marker", "line_cap", "line_join"):
                if stroke.get(key):
                    self.appearance_warning(path + ".stroke." + key, f"Stroke {key} is not exported")
        else:
            shape.line.fill.background()

    def font(self, run: Any, font: Mapping[str, Any], path: str) -> None:
        run.font.name = str(font.get("family") or DEFAULT_FONT["family"])
        run.font.size = Pt(_number(font.get("size", 18), path + ".size", minimum=0.1) * 0.75)
        run.font.color.rgb = _color(font.get("color", "1A2B45"), path + ".color")
        for key in ("bold", "italic", "underline"):
            if font.get(key) is not None:
                if not isinstance(font[key], bool):
                    raise NativeStandardExportError("Font flag must be boolean", path=path + "." + key)
                setattr(run.font, key, font[key])
        properties = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            node = OxmlElement(tag)
            node.set("typeface", run.font.name)
            properties.append(node)
        if font.get("letter_spacing") is not None:
            properties.set("spc", str(round(_number(font["letter_spacing"], path + ".letter_spacing") * 75)))
        if font.get("opacity") is not None:
            _opacity(properties.solidFill, font["opacity"], path + ".opacity")
        if font.get("ellipsis"):
            self.appearance_warning(path + ".ellipsis", "Ellipsis is disabled so the full editable text is preserved")

    def text(self, frame: Any, element: Mapping[str, Any], path: str, *, is_list: bool = False,
             default_font: Mapping[str, Any] | None = None, cell: bool = False) -> None:
        frame.clear()
        frame.word_wrap = True
        # Preserve the original box; Office may shrink text to fit on opening.
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.margin_left = frame.margin_right = _px(6 if cell else 0)
        frame.margin_top = frame.margin_bottom = _px(4 if cell else 0)
        alignment = element.get("alignment") or {}
        if isinstance(alignment, str):
            alignment = {"horizontal": alignment, "vertical": "middle"}
        alignment = _record(alignment, path + ".alignment")
        horizontal = alignment.get("horizontal") or "left"
        vertical = alignment.get("vertical") or ("middle" if cell else "top")
        alignments = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
        anchors = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
        if not isinstance(horizontal, str) or not isinstance(vertical, str) or horizontal not in alignments or vertical not in anchors:
            raise NativeStandardExportError("Unknown text alignment", path=path + ".alignment")
        frame.vertical_anchor = anchors[vertical]
        base_font = {**DEFAULT_FONT, **_font_values(default_font, path + ".font"),
                     **_font_values(element.get("font"), path + ".font")}
        items = element.get("items") if is_list else [element.get("runs", [])]
        if not isinstance(items, list):
            raise NativeStandardExportError("Text items/runs must be arrays", path=path)
        first = True
        for item_index, runs in enumerate(items):
            if not isinstance(runs, list):
                raise NativeStandardExportError("Text runs must be an array", path=path)
            paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            paragraph.alignment = alignments[horizontal]
            paragraph.space_before = Pt(0)
            paragraph.space_after = _px(_number(element.get("gap") or 0, path + ".gap", minimum=0)) if is_list else Pt(0)
            line_height = base_font.get("line_height")
            spacing = _number(line_height, path + ".font.line_height", minimum=0.1) if line_height else 1.15
            if spacing > 2:
                self.reject(path + ".font.line_height", "Line-height >2 is ambiguous between canvas multipliers and pixel heights; use an explicit supported multiplier")
            paragraph.line_spacing = spacing
            if is_list:
                marker = element.get("marker") or "bullet"
                ppr = paragraph._p.get_or_add_pPr()
                if marker == "number":
                    bullet = OxmlElement("a:buAutoNum")
                    bullet.set("type", "arabicPeriod")
                elif marker == "bullet":
                    bullet = OxmlElement("a:buChar")
                    bullet.set("char", "•")
                elif marker == "none":
                    bullet = OxmlElement("a:buNone")
                else:
                    self.reject(path + ".marker", f"List marker {marker!r} is unsupported")
                ppr.append(bullet)
                if marker != "none":
                    indent = float(base_font.get("size", 18)) + _number(element.get("marker_gap") or 6, path + ".marker_gap", minimum=0)
                    ppr.set("marL", str(_px(indent)))
                    ppr.set("indent", str(-_px(indent)))
            for run_index, source_run in enumerate(runs):
                rpath = f"{path}.{'items[' + str(item_index) + '].' if is_list else ''}runs[{run_index}]"
                source_run = _record(source_run, rpath)
                if source_run.get("type") not in (None, "text"):
                    self.reject(rpath, "Only plain text runs are supported; LaTeX is not converted")
                text = source_run.get("text")
                if not isinstance(text, str):
                    raise NativeStandardExportError("Text run must contain a string", path=rpath + ".text")
                if _has_inline_markup(text):
                    self.reject(rpath + ".text", "Inline Markdown is unsupported; resolve it into styled plain text runs")
                font = {**base_font, **_font_values(source_run.get("font"), rpath + ".font")}
                if font.get("line_height") is not None and font.get("line_height") != base_font.get("line_height"):
                    self.appearance_warning(rpath + ".font.line_height", "Per-run line height is not exported; paragraph-level line height is used")
                for part_index, part in enumerate(text.split("\n")):
                    if part_index:
                        paragraph.add_line_break()
                    target = paragraph.add_run()
                    target.text = part
                    self.font(target, font, rpath + ".font")

    def container(self, slide: Any, element: Mapping[str, Any], box: tuple[float, float, float, float],
                  path: str, depth: int) -> None:
        x, y, width, height = box
        radius = element.get("border_radius") or 0
        if isinstance(radius, Mapping):
            radii = [_number(radius.get(key, 0), path + ".border_radius." + key, minimum=0) for key in ("tl", "tr", "bl", "br")]
            if len(set(radii)) != 1:
                self.reject(path + ".border_radius", "Asymmetric corner radii are unsupported")
            radius = radii[0]
        radius = _number(radius, path + ".border_radius", minimum=0)
        if element.get("fill") or element.get("stroke"):
            shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                                          _px(x), _px(y), _px(width), _px(height))
            shape.name = str(element.get("name") or "container")
            self.shape_style(shape, element, path)
            if radius:
                shape.adjustments[0] = min(0.5, radius / min(width, height))
        top, right, bottom, left = _padding(element.get("padding"), path + ".padding")
        inner_width, inner_height = width - left - right, height - top - bottom
        if inner_width <= 0 or inner_height <= 0:
            raise NativeStandardExportError("Padding leaves no content area", path=path)
        kind = element.get("type")
        if kind == "container":
            child = element.get("child")
            if child is not None:
                child = _record(child, path + ".child")
                child_width, child_height = self.bounded_child_size(child, path + ".child", (inner_width, inner_height))
                if child.get("layout") or child.get("__presenton_manual_position") is True:
                    self.reject(path + ".child", "Container child layout/manual-position overrides require resolved geometry")
                alignment = _record(element.get("alignment") or {}, path + ".alignment")
                child_x, child_y = _point(child, path + ".child")
                if child.get("type") != "group":
                    if alignment.get("horizontal") in ("center", "right"):
                        child_x = self.align_offset(alignment.get("horizontal"), inner_width, child_width)
                    if alignment.get("vertical") in ("middle", "bottom"):
                        child_y = self.align_offset(alignment.get("vertical"), inner_height, child_height)
                # Renderer alignment overrides the child's local offset on the
                # centered/end-aligned axis; it does not add both offsets.
                source_x, source_y = _point(child, path + ".child")
                start = len(slide.shapes)
                self.element(slide, child, (x + left + child_x - source_x, y + top + child_y - source_y), path + ".child", depth + 1,
                             (x + left + child_x, y + top + child_y, child_width, child_height))
                for emitted in list(slide.shapes)[start:]:
                    # Canvas containers clip their descendants, whereas native
                    # flattened shapes do not. Refuse content requiring a clip
                    # rather than show material that the source canvas hides.
                    angle = math.radians(emitted.rotation)
                    emitted_w = (abs(math.cos(angle)) * emitted.width + abs(math.sin(angle)) * emitted.height) / EMU_PER_PIXEL
                    emitted_h = (abs(math.sin(angle)) * emitted.width + abs(math.cos(angle)) * emitted.height) / EMU_PER_PIXEL
                    center_x = (emitted.left + emitted.width / 2) / EMU_PER_PIXEL
                    center_y = (emitted.top + emitted.height / 2) / EMU_PER_PIXEL
                    if (center_x - emitted_w / 2 < x - .01 or center_y - emitted_h / 2 < y - .01
                            or center_x + emitted_w / 2 > x + width + .01 or center_y + emitted_h / 2 > y + height + .01):
                        self.reject(path + ".child", "Clipped container descendants are unsupported; resolve visible geometry first")
            return
        children = element.get("children", [])
        if not isinstance(children, list):
            raise NativeStandardExportError("children must be an array", path=path)
        if not children:
            return
        if element.get("wrap"):
            self.reject(path + ".wrap", "Wrapping flex layout is unsupported")
        sizes = []
        for index, child in enumerate(children):
            child = _record(child, f"{path}.children[{index}]")
            if child.get("layout") or child.get("__presenton_manual_position") is True:
                self.reject(f"{path}.children[{index}]", "Flex/grid overrides, spans and manual positions are unsupported")
            if not isinstance(child.get("type"), str):
                raise NativeStandardExportError("Element type must be a string", path=f"{path}.children[{index}].type")
            if child.get("type") == "vector":
                self.reject(f"{path}.children[{index}]", "Vectors inside flow layouts require resolved absolute geometry")
            sizes.append(_size(child, f"{path}.children[{index}]"))
        gap = _number(element.get("gap") or 0, path + ".gap", minimum=0)
        if kind == "flex":
            direction = element.get("direction") or "column"
            if not isinstance(direction, str) or direction not in {"row", "column"}:
                raise NativeStandardExportError("Unknown flex direction", path=path)
            column = direction == "column"
            gap = _number(element.get("row_gap" if column else "column_gap") or gap, path + ".gap", minimum=0)
            used = sum(size[1 if column else 0] for size in sizes) + gap * (len(sizes) - 1)
            available = inner_height if column else inner_width
            if used > available + 0.01:
                self.reject(path, "Flex shrinking is unsupported; resolve layout geometry before export")
            cursor = self.align_offset(element.get("justify_content"), available, used)
            align = element.get("align_items") or "stretch"
            for index, (child, (child_width, child_height)) in enumerate(zip(children, sizes)):
                cross = inner_width if column else inner_height
                cross_size = child_width if column else child_height
                if align == "stretch":
                    cross_size = cross
                cross_offset = self.align_offset(align, cross, cross_size)
                child_box = (x + left + (cross_offset if column else cursor),
                             y + top + (cursor if column else cross_offset),
                             cross_size if column else child_width,
                             child_height if column else cross_size)
                self.element(slide, child, (0, 0), f"{path}.children[{index}]", depth + 1, child_box)
                cursor += (child_height if column else child_width) + gap
        else:
            columns = _number(element.get("columns"), path + ".columns", minimum=1)
            rows = _number(element.get("rows") or math.ceil(len(children) / columns), path + ".rows", minimum=1)
            if columns != int(columns) or rows != int(rows) or rows * columns < len(children):
                raise NativeStandardExportError("Grid dimensions must be positive integers covering all children", path=path)
            column_gap = _number(element.get("column_gap") or gap, path + ".column_gap", minimum=0)
            row_gap = _number(element.get("row_gap") or gap, path + ".row_gap", minimum=0)
            cell_width = (inner_width - (columns - 1) * column_gap) / columns
            cell_height = (inner_height - (rows - 1) * row_gap) / rows
            if cell_width <= 0 or cell_height <= 0:
                raise NativeStandardExportError("Grid gaps leave no content area", path=path)
            for index, (child, (child_width, child_height)) in enumerate(zip(children, sizes)):
                horizontal = element.get("justify_items") or "stretch"
                vertical = element.get("align_items") or "stretch"
                child_width = cell_width if horizontal == "stretch" else child_width
                child_height = cell_height if vertical == "stretch" else child_height
                child_box = (x + left + (index % int(columns)) * (cell_width + column_gap) + self.align_offset(horizontal, cell_width, child_width),
                             y + top + (index // int(columns)) * (cell_height + row_gap) + self.align_offset(vertical, cell_height, child_height),
                             child_width, child_height)
                self.element(slide, child, (0, 0), f"{path}.children[{index}]", depth + 1, child_box)

    @staticmethod
    def align_offset(value: Any, available: float, used: float) -> float:
        if value in (None, "left", "top", "start", "flex-start", "stretch"):
            return 0
        if value in ("middle", "center"):
            return max(0, available - used) / 2
        if value in ("right", "bottom", "end", "flex-end"):
            return max(0, available - used)
        raise NativeStandardExportError(f"Unsupported alignment {value!r}", path="alignment")

    def vector(self, slide: Any, element: Mapping[str, Any], offset: tuple[float, float], path: str) -> None:
        radii = element.get("corner_radii")
        if radii is not None and not isinstance(radii, list):
            raise NativeStandardExportError("corner_radii must be an array", path=path + ".corner_radii")
        parsed_radii = [_number(value, f"{path}.corner_radii[{index}]", minimum=0) for index, value in enumerate(radii or [])]
        if element.get("curve"):
            self.reject(path, "Curved polygon vectors are unsupported")
        for marker in ("start_marker", "end_marker"):
            if element.get(marker) not in (None, "none"):
                self.reject(path + "." + marker, "Vector markers are unsupported")
        points = element.get("points")
        if not isinstance(points, list) or len(points) < 2:
            raise NativeStandardExportError("Vector requires at least two points", path=path + ".points")
        pairs = []
        for index, point in enumerate(points):
            point = _record(point, f"{path}.points[{index}]")
            pairs.append((_number(point.get("x"), path + ".points.x") + offset[0],
                          _number(point.get("y"), path + ".points.y") + offset[1]))
        shape_type = element.get("shape") or "polygon"
        if shape_type == "ellipse":
            xs, ys = zip(*pairs)
            if max(xs) <= min(xs) or max(ys) <= min(ys):
                raise NativeStandardExportError("Ellipse bounds must have positive dimensions", path=path)
            shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, _px(min(xs)), _px(min(ys)),
                                          _px(max(xs) - min(xs)), _px(max(ys) - min(ys)))
        elif shape_type == "polygon":
            closed = len(pairs) > 2 if element.get("closed") is None else element["closed"]
            if not isinstance(closed, bool):
                raise NativeStandardExportError("closed must be a boolean", path=path + ".closed")
            if any(parsed_radii):
                xs, ys = zip(*pairs)
                corners = {(min(xs), min(ys)), (max(xs), min(ys)), (max(xs), max(ys)), (min(xs), max(ys))}
                is_rectangle = len(pairs) == 4 and set(pairs) == corners and all(
                    pairs[index][0] == pairs[(index + 1) % 4][0] or pairs[index][1] == pairs[(index + 1) % 4][1] for index in range(4))
                if not closed or not is_rectangle or len(parsed_radii) != 4 or len(set(parsed_radii)) != 1:
                    self.reject(path + ".corner_radii", "Only uniform rounded axis-aligned rectangular vectors are supported")
                vector_width, vector_height = max(xs) - min(xs), max(ys) - min(ys)
                if min(vector_width, vector_height) <= 0:
                    raise NativeStandardExportError("Rounded rectangle needs positive dimensions", path=path)
                shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, _px(min(xs)), _px(min(ys)), _px(vector_width), _px(vector_height))
                shape.adjustments[0] = min(.5, parsed_radii[0] / min(vector_width, vector_height))
            elif len(pairs) == 2 and not closed:
                shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *[_px(value) for pair in pairs for value in pair])
            else:
                builder = slide.shapes.build_freeform(*pairs[0], scale=EMU_PER_PIXEL)
                builder.add_line_segments(pairs[1:], close=closed)
                shape = builder.convert_to_shape()
        else:
            self.reject(path + ".shape", f"Vector shape {shape_type!r} is unsupported")
        if element.get("rotation"):
            shape.rotation = _number(element["rotation"], path + ".rotation")
        if element.get("flip_h") or element.get("flip_v"):
            transform = shape._element.spPr.get_or_add_xfrm()
            if element.get("flip_h"):
                transform.set("flipH", "1")
            if element.get("flip_v"):
                transform.set("flipV", "1")
        if hasattr(shape, "fill"):
            self.shape_style(shape, element, path)
        else:
            stroke = element.get("stroke") or {}
            shape.line.color.rgb = _color(stroke.get("color", "1A2B45"), path + ".stroke.color")
            shape.line.width = _px(_number(stroke.get("width", 1), path + ".stroke.width", minimum=0))
        shape.name = str(element.get("name") or "vector")

    def table(self, slide: Any, element: Mapping[str, Any], box: tuple[float, float, float, float], path: str) -> Any:
        headers, rows = element.get("columns"), element.get("rows")
        if not isinstance(headers, list) or not headers or not isinstance(rows, list):
            raise NativeStandardExportError("Table needs columns and rows arrays", path=path)
        if any(not isinstance(row, list) or len(row) != len(headers) for row in rows):
            raise NativeStandardExportError("Table rows must match the column count", path=path + ".rows")
        if len(headers) > 100 or len(rows) > 1000:
            raise NativeStandardExportError("Table exceeds export limits", path=path)
        x, y, width, height = box
        shape = slide.shapes.add_table(len(rows) + 1, len(headers), _px(x), _px(y), _px(width), _px(height))
        table = shape.table
        table.first_row = False
        table.horz_banding = False
        # Strip the presentation theme's table style: source cell colors win.
        table_style = table._tbl.tblPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId")
        if table_style is not None:
            table._tbl.tblPr.remove(table_style)
        for row_index, row in enumerate([headers, *rows]):
            for col_index, source_cell in enumerate(row):
                cpath = f"{path}.{'columns' if row_index == 0 else 'rows[' + str(row_index - 1) + ']'}[{col_index}]"
                source_cell = _record(source_cell, cpath)
                cell = table.cell(row_index, col_index)
                fill = source_cell.get("fill") or source_cell.get("color")
                if fill:
                    fill = _record(fill, cpath + ".color")
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _color(fill.get("color"), cpath + ".color.color")
                    if fill.get("opacity") is not None:
                        _opacity(cell.fill._xPr.solidFill, fill["opacity"], cpath + ".color.opacity")
                else:
                    cell.fill.background()
                default_font = _font_values(element.get("font"), path + ".font")
                cell_runs = source_cell.get("runs", [])
                if not isinstance(cell_runs, list):
                    raise NativeStandardExportError("Text runs must be an array", path=cpath + ".runs")
                first_run = _record((cell_runs or [{}])[0], cpath + ".runs[0]")
                first_run_font = _font_values(first_run.get("font"), cpath + ".runs[0].font")
                if row_index == 0 and "bold" not in _font_values(source_cell.get("font"), cpath + ".font") and "bold" not in first_run_font:
                    default_font["bold"] = True
                self.text(cell.text_frame, source_cell, cpath, default_font=default_font, cell=True)
                borders = source_cell.get("borders")
                if borders is not None:
                    borders = _record(borders, cpath + ".borders")
                for side, xml_name in (("left", "lnL"), ("right", "lnR"), ("top", "lnT"), ("bottom", "lnB")):
                    border = (borders.get(side) if borders else (source_cell.get("stroke") or {"color": "D0D5DD", "width": 1}))
                    if border is not None:
                        border = _record(border, cpath + ".borders." + side)
                    tcpr = cell._tc.get_or_add_tcPr()
                    line = OxmlElement("a:" + xml_name)
                    if border and _number(border.get("width", 0), cpath + ".borders." + side + ".width", minimum=0) > 0:
                        line.set("w", str(_px(_number(border["width"], cpath + ".borders.width", minimum=0))))
                        solid = OxmlElement("a:solidFill")
                        color = OxmlElement("a:srgbClr")
                        color.set("val", str(_color(border.get("color"), cpath + ".borders.color")))
                        solid.append(color)
                        line.append(solid)
                        if border.get("opacity") is not None:
                            _opacity(solid, border["opacity"], cpath + ".borders.opacity")
                        if border.get("dash"):
                            self.appearance_warning(cpath + ".borders", "Dashed table borders are exported as solid lines")
                    else:
                        line.append(OxmlElement("a:noFill"))
                    tcpr.append(line)
        return shape

    def chart(self, slide: Any, element: Mapping[str, Any], box: tuple[float, float, float, float], path: str) -> Any:
        kind = element.get("chart_type") or element.get("chartType")
        types = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
                 "line": XL_CHART_TYPE.LINE, "donut": XL_CHART_TYPE.DOUGHNUT, "doughnut": XL_CHART_TYPE.DOUGHNUT}
        if not isinstance(kind, str) or kind not in types:
            self.reject(path + ".chart_type", f"Chart type {kind!r} is unsupported")
        categories, series = element.get("categories"), element.get("series")
        if not isinstance(categories, list) or not categories or not all(isinstance(value, str) for value in categories):
            raise NativeStandardExportError("Chart categories must be a non-empty string array", path=path + ".categories")
        if not isinstance(series, list) or not series or len(series) > 50 or len(categories) > 10000:
            raise NativeStandardExportError("Chart needs 1–50 series and at most 10000 categories", path=path + ".series")
        donut = kind in {"donut", "doughnut"}
        if donut and len(series) != 1:
            self.reject(path + ".series", "Doughnut charts require exactly one series; extra series will not be discarded")
        for index, category in enumerate(categories):
            if len(category) > MAX_WORKBOOK_STRING_LENGTH:
                raise NativeStandardExportError("Category label exceeds the XLSX string-cell limit (32767 characters)", path=f"{path}.categories[{index}]", code="export_limit")
            if _has_inline_markup(category):
                self.reject(f"{path}.categories[{index}]", "Chart label Markdown is unsupported; use plain text labels")
        data = _LiteralChartData()
        data.categories = categories
        for index, item in enumerate(series):
            item = _record(item, f"{path}.series[{index}]")
            name, values = item.get("name"), item.get("values")
            if not isinstance(name, str) or not isinstance(values, list) or len(values) != len(categories):
                raise NativeStandardExportError("Each series requires a name and one value per category", path=f"{path}.series[{index}]")
            if len(name) > MAX_WORKBOOK_STRING_LENGTH:
                raise NativeStandardExportError("Series name exceeds the XLSX string-cell limit (32767 characters)", path=f"{path}.series[{index}].name", code="export_limit")
            if _has_inline_markup(name):
                self.reject(f"{path}.series[{index}].name", "Series-name Markdown is unsupported; use plain text names")
            parsed = [_number(value, f"{path}.series[{index}].values[{j}]", minimum=0 if donut else None) for j, value in enumerate(values)]
            if donut and sum(parsed) <= 0:
                raise NativeStandardExportError("Doughnut chart requires a positive total", path=path)
            data.add_series(name, parsed)
        shape = slide.shapes.add_chart(types[kind], *[_px(value) for value in box], data)
        chart = shape.chart
        chart.font.name = "Arial"
        chart.font.size = Pt(13.5)
        chart.font.color.rgb = _color(element.get("text_color") or "1A2B45", path + ".text_color")
        auto_legend = donut or len(series) > 1 or bool(series[0].get("name") and series[0].get("name") != "Series 1")
        chart.has_legend = bool(auto_legend if element.get("legend") is None else element["legend"])
        if chart.has_legend:
            positions = {"bottom": XL_LEGEND_POSITION.BOTTOM, "top": XL_LEGEND_POSITION.TOP,
                         "left": XL_LEGEND_POSITION.LEFT, "right": XL_LEGEND_POSITION.RIGHT}
            position = element.get("legend_position") or "bottom"
            if not isinstance(position, str) or position not in positions:
                raise NativeStandardExportError("Unknown legend position", path=path + ".legend_position")
            chart.legend.position = positions[position]
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(12)
            chart.legend.font.color.rgb = _color(element.get("legend_color") or element.get("text_color") or "1A2B45", path + ".legend_color")
        title = element.get("title")
        chart.has_title = bool(title)
        if title:
            if not isinstance(title, str):
                raise NativeStandardExportError("Chart title must be a string", path=path + ".title")
            self.text(chart.chart_title.text_frame, {"runs": [{"text": title}], "font": {"size": 22, "bold": True, "color": element.get("title_color") or "1A2B45"}}, path + ".title")
        colors = element.get("colors") or ([element["color"]] if element.get("color") else DEFAULT_CHART_COLORS)
        if not isinstance(colors, (list, tuple)) or not colors:
            raise NativeStandardExportError("Chart colors must be a non-empty array", path=path + ".colors")
        palette = [_color(color, path + ".colors") for color in colors]
        for index, target in enumerate(chart.series):
            if kind == "line":
                target.format.line.color.rgb = palette[index % len(palette)]
                target.format.line.width = Pt(2)
                target.smooth = False
            else:
                target.format.fill.solid()
                target.format.fill.fore_color.rgb = palette[index % len(palette)]
                target.format.line.fill.background()
                if len(series) == 1:
                    for point_index, point in enumerate(target.points):
                        point.format.fill.solid()
                        point.format.fill.fore_color.rgb = palette[point_index % len(palette)]
                        point.format.line.fill.background()
        plot = chart.plots[0]
        label = element.get("data_labels")
        if label is True:
            label = "top"
        plot.has_data_labels = bool(label)
        if label:
            labels = {"base": XL_DATA_LABEL_POSITION.INSIDE_BASE, "mid": XL_DATA_LABEL_POSITION.CENTER,
                      "top": XL_DATA_LABEL_POSITION.ABOVE if kind == "line" else XL_DATA_LABEL_POSITION.INSIDE_END,
                      "outside": XL_DATA_LABEL_POSITION.ABOVE if kind == "line" else XL_DATA_LABEL_POSITION.OUTSIDE_END}
            if not isinstance(label, str) or label not in labels:
                raise NativeStandardExportError("Unknown data label position", path=path + ".data_labels")
            plot.data_labels.position = labels[label]
            plot.data_labels.show_value = True
            plot.data_labels.show_legend_key = False
            plot.data_labels.font.size = Pt(11)
            plot.data_labels.font.color.rgb = _color(element.get("text_color") or "1A2B45", path + ".text_color")
        if donut:
            plot.hole_size = 60
        else:
            for axis_name, axis in (("x", chart.category_axis), ("y", chart.value_axis)):
                axis.visible = element.get(axis_name + "_axis", True) is not False
                axis.has_major_gridlines = element.get(axis_name + "_axis_grid") is not False
                axis.format.line.color.rgb = _color(element.get("axis_color") or "D0D5DD", path + ".axis_color")
                axis.tick_labels.font.size = Pt(11)
                if axis.has_major_gridlines:
                    axis.major_gridlines.format.line.color.rgb = _color(element.get("grid_color") or "E5E7EB", path + ".grid_color")
                axis_title = element.get(axis_name + "_axis_title")
                axis.has_title = bool(axis_title)
                if axis_title:
                    self.text(axis.axis_title.text_frame, {"runs": [{"text": axis_title}], "font": {"size": 16}}, path + "." + axis_name + "_axis_title")
        if element.get("source"):
            self.appearance_warning(path + ".source", "Chart source metadata is not displayed; add it as a text element to show it")
        return shape


def export_standard_pptx(
    slide_uis: Sequence[Mapping[str, Any]], *, title: str = "Presentation",
    theme: Mapping[str, Any] | None = None,
    image_assets: Mapping[str, bytes] | None = None,
    unsupported: Literal["error", "warn"] = "error",
) -> StandardPptxExport:
    """Return PPTX bytes and warnings without writing files or changing inputs.

    Pass the exact already-resolved ``ui`` for each slide, in display order.
    The caller owns authorization, presentation version checks, persistence,
    filename selection, authorized image-byte resolution, and exposing warnings
    to the user. image_assets keys match image.data exactly; values are PNG/JPEG
    bytes only, never file paths, URLs, streams, or callbacks. Never treat this
    exporter as a successful substitute for an unsupported presentation mode.
    """
    if not isinstance(unsupported, str) or unsupported not in {"error", "warn"}:
        raise ValueError("unsupported must be 'error' or 'warn'")
    if not isinstance(slide_uis, (list, tuple)) or not 1 <= len(slide_uis) <= MAX_SLIDES:
        raise NativeStandardExportError(f"Expected 1–{MAX_SLIDES} slide UI objects", path="slides")
    if not isinstance(title, str):
        raise NativeStandardExportError("Title must be a string", path="title")
    if theme is not None:
        _record(theme, "theme")
        data = _record(theme.get("data", theme), "theme.data")
        if "colors" in data:
            _record(data["colors"], "theme.data.colors")
    _validate_complexity(slide_uis)
    validated_images = _validate_image_assets(image_assets)
    return _Exporter(title, theme, unsupported, validated_images).run(slide_uis)
