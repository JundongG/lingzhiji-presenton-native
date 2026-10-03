# SPDX-License-Identifier: Apache-2.0
# Independent native-export add-on; no licensed export-core implementation is used.
"""Offline structural and native-edit tests; also runnable without pytest.

Run: python -m unittest discover -s tests/unit -p test_native_standard_pptx.py
These checks do not claim that the PowerPoint desktop UI was exercised.
"""

import copy
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from pptx import Presentation
from pptx.chart.data import CategoryChartData

from services.native_standard_pptx import (
    EMU_PER_PIXEL,
    NativeStandardExportError,
    export_standard_pptx,
)

FIXTURE_PATH = Path(__file__).parents[1] / "fixtures" / "native_standard_zh.json"
NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
}


def fixture(revised=False):
    value = json.loads(FIXTURE_PATH.read_text(encoding="utf-8"))
    if revised:
        revision = value["revision"]
        slides = value["slides"]
        slides[0]["components"][0]["elements"][0]["runs"][0]["text"] = revision["cover_title"]
        slides[1]["components"][1]["elements"][0]["rows"][0][2]["runs"][0]["text"] = revision["table_budget"]
        for slide_index, key, series_index in ((2, "column_values", 0), (3, "line_values", 1), (4, "donut_values", 0)):
            slides[slide_index]["components"][1]["elements"][0]["series"][series_index]["values"] = revision[key]
        value["title"] += "（修订版）"
    return value


def workbook_rows(blob):
    """Read raw embedded XLSX data using stdlib, without a new dependency."""
    with ZipFile(BytesIO(blob)) as archive:
        strings = []
        if "xl/sharedStrings.xml" in archive.namelist():
            strings = ["".join(item.itertext()) for item in ET.fromstring(archive.read("xl/sharedStrings.xml"))]
        tree = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        rows = []
        for row in tree.findall("s:sheetData/s:row", NS):
            values = []
            for cell in row.findall("s:c", NS):
                value = cell.find("s:v", NS)
                if cell.get("t") == "s":
                    values.append(strings[int(value.text)])
                elif cell.get("t") == "inlineStr":
                    values.append("".join(cell.find("s:is", NS).itertext()))
                elif value is not None:
                    values.append(float(value.text))
                else:
                    values.append(None)
            rows.append(values)
        return rows


def all_workbook_rows(pptx_blob):
    with ZipFile(BytesIO(pptx_blob)) as archive:
        return [workbook_rows(archive.read(name)) for name in sorted(archive.namelist())
                if name.startswith("ppt/embeddings/") and name.endswith(".xlsx")]


def edited_native_roundtrip(blob):
    """Exercise python-pptx's native table/chart editing API, never image editing."""
    presentation = Presentation(BytesIO(blob))
    table = next(shape.table for shape in presentation.slides[1].shapes if shape.has_table)
    # Edit the existing run rather than assigning cell.text, which rebuilds
    # the text frame and discards its original run font/color formatting.
    table.cell(1, 2).text_frame.paragraphs[0].runs[0].text = "177（原生二次修改）"
    chart = next(shape.chart for shape in presentation.slides[2].shapes if shape.has_chart)
    data = CategoryChartData()
    data.categories = ["第一季度", "第二季度", "第三季度", "第四季度"]
    data.add_series("二次修改后的完成项目", [19, 26, 35, 48])
    chart.replace_data(data)
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


class NativeStandardPptxTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.input = fixture()
        cls.result = export_standard_pptx(cls.input["slides"], title=cls.input["title"])

    def test_five_slide_fixture_is_fully_native_and_has_workbooks(self):
        self.assertEqual(self.result.warnings, ())
        with ZipFile(BytesIO(self.result.pptx_bytes)) as archive:
            slides = [ET.fromstring(archive.read(f"ppt/slides/slide{i}.xml")) for i in range(1, 6)]
            self.assertGreaterEqual(sum(len(slide.findall(".//p:sp", NS)) for slide in slides), 20)
            self.assertEqual(sum(len(slide.findall(".//a:tbl", NS)) for slide in slides), 1)
            self.assertEqual(sum(len(slide.findall(".//c:chart", NS)) for slide in slides), 3)
            self.assertEqual(sum(len(slide.findall(".//p:pic", NS)) for slide in slides), 0)
            charts = [name for name in archive.namelist() if name.startswith("ppt/charts/chart") and name.endswith(".xml")]
            self.assertEqual(len(charts), 3)
            for name, kind in zip(sorted(charts), ("barChart", "lineChart", "doughnutChart")):
                chart = ET.fromstring(archive.read(name))
                self.assertIsNotNone(chart.find(".//c:" + kind, NS))
                self.assertIsNotNone(chart.find(".//c:externalData", NS))
            self.assertFalse(any(name.startswith("ppt/media/") for name in archive.namelist()))
        workbooks = all_workbook_rows(self.result.pptx_bytes)
        self.assertEqual(len(workbooks), 3)
        self.assertEqual(workbooks[0][1:], [["第一季度", 12], ["第二季度", 18], ["第三季度", 25], ["第四季度", 32]])
        self.assertEqual(workbooks[1][-1], ["十月", 80, 84])
        self.assertEqual(workbooks[2][1:], [["研发", 50], ["验证", 30], ["交付", 20]])

    def test_geometry_font_and_chinese_text_survive(self):
        presentation = Presentation(BytesIO(self.result.pptx_bytes))
        self.assertEqual(len(presentation.slides), 5)
        self.assertEqual(presentation.slide_width, 1280 * EMU_PER_PIXEL)
        self.assertEqual(presentation.slide_height, 720 * EMU_PER_PIXEL)
        title = next(shape for shape in presentation.slides[0].shapes if shape.name == "标题")
        self.assertEqual((title.left, title.top, title.width, title.height),
                         tuple(value * EMU_PER_PIXEL for value in (64, 44, 1152, 96)))
        self.assertEqual(title.text, self.input["slides"][0]["components"][0]["elements"][0]["runs"][0]["text"])
        self.assertEqual(title.text_frame.paragraphs[0].runs[0].font.name, "Noto Sans CJK SC")
        self.assertEqual(title.text_frame.paragraphs[0].runs[0].font.size.pt, 27)
        self.assertIn("Noto Sans CJK SC", title._element.xml)
        table = next(shape for shape in presentation.slides[1].shapes if shape.has_table)
        self.assertEqual((table.left, table.top, table.width, table.height),
                         tuple(value * EMU_PER_PIXEL for value in (64, 220, 1152, 340)))
        self.assertEqual(table.table.cell(0, 0).text, "项目")
        self.assertEqual(table.table.cell(1, 2).text, "150")

    def test_revision_preserves_long_title_and_all_changed_chart_values(self):
        value = fixture(revised=True)
        revised = export_standard_pptx(value["slides"], title=value["title"])
        presentation = Presentation(BytesIO(revised.pptx_bytes))
        title = next(shape for shape in presentation.slides[0].shapes if shape.name == "标题")
        self.assertEqual(title.text, value["revision"]["cover_title"])
        self.assertEqual(title.height, 96 * EMU_PER_PIXEL)
        self.assertIn("normAutofit", title._element.xml)
        table = next(shape.table for shape in presentation.slides[1].shapes if shape.has_table)
        self.assertEqual(table.cell(1, 2).text, "168")
        for workbook, expected, column in zip(all_workbook_rows(revised.pptx_bytes),
                                             ([15, 22, 29, 41], [50, 59, 73, 82, 91], [46, 34, 20]), (1, 2, 1)):
            self.assertEqual([row[column] for row in workbook[1:]], expected)

    def test_native_table_and_chart_edit_resave_roundtrip(self):
        blob = edited_native_roundtrip(self.result.pptx_bytes)
        presentation = Presentation(BytesIO(blob))
        table = next(shape.table for shape in presentation.slides[1].shapes if shape.has_table)
        self.assertEqual(table.cell(1, 2).text, "177（原生二次修改）")
        edited_font = table.cell(1, 2).text_frame.paragraphs[0].runs[0].font
        original_table = next(shape.table for shape in Presentation(BytesIO(self.result.pptx_bytes)).slides[1].shapes if shape.has_table)
        original_font = original_table.cell(1, 2).text_frame.paragraphs[0].runs[0].font
        self.assertEqual((edited_font.name, edited_font.size, edited_font.bold, str(edited_font.color.rgb)),
                         (original_font.name, original_font.size, original_font.bold, str(original_font.color.rgb)))
        chart = next(shape.chart for shape in presentation.slides[2].shapes if shape.has_chart)
        self.assertEqual(chart.series[0].name, "二次修改后的完成项目")
        self.assertEqual(list(chart.series[0].values), [19, 26, 35, 48])
        self.assertEqual(all_workbook_rows(blob)[0][1:], [["第一季度", 19], ["第二季度", 26], ["第三季度", 35], ["第四季度", 48]])

    def test_input_is_not_mutated_and_no_network_or_temporary_files(self):
        value = fixture()
        original = copy.deepcopy(value)
        with patch("socket.socket", side_effect=AssertionError("Network forbidden")), \
             patch.object(tempfile, "mkstemp", side_effect=AssertionError("Temp files forbidden")):
            export_standard_pptx(value["slides"], title=value["title"])
        self.assertEqual(value, original)

    def test_unknown_type_raises_with_exact_path(self):
        with self.assertRaises(NativeStandardExportError) as caught:
            export_standard_pptx([{"components": [{"position": {"x": 0, "y": 0}, "elements": [{"type": "video", "data": "https://example.invalid/x.png"}]}]}])
        self.assertEqual(caught.exception.path, "slides[0].ui.components[0].elements[0]")
        self.assertEqual(caught.exception.code, "unsupported_element")

    def test_warn_mode_omits_unsupported_content_and_reports_it(self):
        ui = {"elements": [{"type": "video", "data": "never-fetch-this"}, self.simple_text()]}
        result = export_standard_pptx([ui], unsupported="warn")
        self.assertEqual(len(result.warnings), 1)
        self.assertIn("omitted", result.warnings[0].message)
        self.assertEqual(result.warnings[0].path, "slides[0].ui.elements[0]")
        self.assertEqual(len(Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes), 1)

    def test_latex_never_silently_rasterizes_or_discards_a_run(self):
        element = self.simple_text()
        element["runs"] = [{"text": "Prefix"}, {"type": "latex", "latex": "x^2"}]
        with self.assertRaisesRegex(NativeStandardExportError, "LaTeX"):
            export_standard_pptx([{"elements": [element]}])
        result = export_standard_pptx([{"elements": [element]}], unsupported="warn")
        self.assertEqual(len(Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes), 0)
        self.assertIn("runs[1]", result.warnings[0].path)

    def test_invalid_chart_data_is_rejected_without_zero_filling(self):
        base = fixture()["slides"][2]["components"][1]["elements"][0]
        for value in ([1, 2], [1, 2, float("nan"), 4], [1, 2, True, 4], [1, 2, "3", 4]):
            chart = copy.deepcopy(base)
            chart["series"][0]["values"] = value
            with self.subTest(value=value), self.assertRaises(NativeStandardExportError):
                export_standard_pptx([{"elements": [chart]}], unsupported="warn")

    def test_doughnut_does_not_discard_additional_series(self):
        chart = fixture()["slides"][4]["components"][1]["elements"][0]
        chart["series"].append({"name": "不可丢失", "values": [1, 2, 3]})
        with self.assertRaisesRegex(NativeStandardExportError, "extra series will not be discarded"):
            export_standard_pptx([{"elements": [chart]}])

    def test_table_shape_must_be_rectangular(self):
        table = fixture()["slides"][1]["components"][1]["elements"][0]
        table["rows"][0].pop()
        with self.assertRaisesRegex(NativeStandardExportError, "column count"):
            export_standard_pptx([{"elements": [table]}])

    def test_nested_offsets_and_container_alignment(self):
        text = self.simple_text()
        text["position"] = {"x": 5, "y": 7}
        ui = {"components": [{"position": {"x": 10, "y": 20}, "elements": [
            {"type": "group", "position": {"x": 30, "y": 40}, "children": [
                {"type": "container", "position": {"x": 50, "y": 60}, "size": {"width": 300, "height": 200},
                 "padding": 10, "alignment": {"horizontal": "center", "vertical": "middle"}, "child": text}
            ]}]}]}
        result = export_standard_pptx([ui])
        shape = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes[0]
        # inner 280x180, text 100x40; centering replaces the local offset.
        self.assertEqual((shape.left, shape.top), (190 * EMU_PER_PIXEL, 200 * EMU_PER_PIXEL))

    def test_simple_fixed_flex_layout(self):
        ui = {"elements": [{"type": "flex", "position": {"x": 100, "y": 100},
                            "size": {"width": 300, "height": 100}, "direction": "row", "gap": 20,
                            "align_items": "center", "justify_content": "center",
                            "children": [self.simple_text(), self.simple_text()]}]}
        result = export_standard_pptx([ui])
        shapes = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes
        self.assertEqual([(shape.left, shape.top) for shape in shapes],
                         [(140 * EMU_PER_PIXEL, 130 * EMU_PER_PIXEL), (260 * EMU_PER_PIXEL, 130 * EMU_PER_PIXEL)])

    def test_equal_cell_grid_layout(self):
        ui = {"elements": [{"type": "grid", "position": {"x": 100, "y": 100}, "size": {"width": 220, "height": 100},
                            "columns": 2, "gap": 20, "children": [self.simple_text(), self.simple_text()]}]}
        result = export_standard_pptx([ui])
        shapes = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes
        self.assertEqual([(shape.left, shape.top, shape.width, shape.height) for shape in shapes],
                         [(100 * EMU_PER_PIXEL, 100 * EMU_PER_PIXEL, 100 * EMU_PER_PIXEL, 100 * EMU_PER_PIXEL),
                          (220 * EMU_PER_PIXEL, 100 * EMU_PER_PIXEL, 100 * EMU_PER_PIXEL, 100 * EMU_PER_PIXEL)])

    def test_unresolved_and_advanced_geometry_is_explicitly_rejected(self):
        bad = [
            {"type": "text", "runs": [{"text": "No geometry"}]},
            {"type": "group", "rotation": 20, "children": [self.simple_text()]},
            {"type": "flex", "wrap": True, "size": {"width": 200, "height": 100}, "children": [self.simple_text()]},
            {"type": "vector", "shape": "polygon", "curve": {"type": "smooth"}, "points": [{"x": 0, "y": 0}, {"x": 20, "y": 20}]},
        ]
        for element in bad:
            with self.subTest(element=element), self.assertRaises(NativeStandardExportError):
                export_standard_pptx([{"elements": [element]}])

    def test_basic_polygons_strokes_and_fill_alpha_are_native(self):
        ui = {"elements": [{"type": "vector", "closed": True, "points": [{"x": 10, "y": 10}, {"x": 100, "y": 10}, {"x": 50, "y": 80}],
                            "fill": {"color": "#ABC", "opacity": .5}, "stroke": {"color": "#123456", "width": 2, "opacity": .75}}]}
        result = export_standard_pptx([ui])
        xml = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes[0]._element.xml
        self.assertIn("custGeom", xml)
        self.assertIn('val="AABBCC"', xml)
        self.assertIn('val="50000"', xml)
        self.assertIn('val="75000"', xml)

    def test_out_of_bounds_geometry_is_preserved_with_warning(self):
        element = self.simple_text()
        element["position"] = {"x": -10, "y": 0}
        result = export_standard_pptx([{"elements": [element]}])
        self.assertEqual(result.warnings[0].code, "appearance_difference")
        self.assertEqual(Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes[0].left, -10 * EMU_PER_PIXEL)

    def test_fixture_matches_upstream_current_schema_and_validator(self):
        from templates.v2.models.layouts import SlideLayout
        from services.chat.slide_ui_helpers import _validate_visual_insert_tree

        for revised in (False, True):
            value = fixture(revised=revised)
            for ui in value["slides"]:
                validated = SlideLayout.model_validate(ui)
                _validate_visual_insert_tree(ui)
                # model_dump retains schema null defaults; these must not break
                # font inheritance or turn null booleans into false values.
                result = export_standard_pptx([validated.model_dump(mode="json")])
                self.assertFalse(result.warnings)

    def test_null_font_properties_inherit_and_run_overrides_survive(self):
        element = self.simple_text()
        element["font"] = {"size": 24, "family": "Noto Sans CJK SC", "color": "#112233", "bold": True}
        element["runs"] = [{"text": "继承", "font": {"size": None, "family": None, "bold": None}},
                           {"text": "覆盖", "font": {"bold": False, "size": 20}}]
        result = export_standard_pptx([{"elements": [element]}])
        runs = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes[0].text_frame.paragraphs[0].runs
        self.assertEqual(runs[0].font.size.pt, 18)
        self.assertTrue(runs[0].font.bold)
        self.assertEqual(runs[0].font.name, "Noto Sans CJK SC")
        self.assertEqual(runs[1].font.size.pt, 15)
        self.assertFalse(runs[1].font.bold)

    def test_source_without_shadows_does_not_acquire_office_theme_effects(self):
        presentation = Presentation(BytesIO(self.result.pptx_bytes))
        for slide in presentation.slides:
            for shape in slide.shapes:
                if shape.shape_type == 1:  # MSO_SHAPE_TYPE.AUTO_SHAPE
                    self.assertTrue(shape._element.xpath("./p:spPr/a:effectLst"))
                    for effect_ref in shape._element.xpath("./p:style/a:effectRef"):
                        self.assertEqual(effect_ref.get("idx"), "0")
                    self.assertFalse(shape._element.xpath(".//a:outerShdw"))

    def test_line_fixture_avoids_dense_labels_without_losing_data(self):
        element = self.input["slides"][3]["components"][1]["elements"][0]
        self.assertIsNone(element["data_labels"])
        presentation = Presentation(BytesIO(self.result.pptx_bytes))
        chart = next(shape.chart for shape in presentation.slides[3].shapes if shape.has_chart)
        self.assertFalse(chart.plots[0].has_data_labels)
        self.assertEqual(len(chart.series), 2)
        self.assertTrue(chart.has_legend)

    def test_background_images_are_never_silently_omitted(self):
        ui = {"background_image": {"data": "https://example.invalid/no-fetch"}, "elements": [self.simple_text()]}
        with self.assertRaisesRegex(NativeStandardExportError, "Background images"):
            export_standard_pptx([ui])
        result = export_standard_pptx([ui], unsupported="warn")
        self.assertEqual(result.warnings[0].path, "slides[0].ui.background_image")
        self.assertIn("omitted", result.warnings[0].message)
        self.assertEqual(len(Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes), 1)

    def test_element_opacity_is_not_silently_ignored(self):
        for opacity in (0, .5):
            element = self.simple_text()
            element["opacity"] = opacity
            with self.subTest(opacity=opacity), self.assertRaisesRegex(NativeStandardExportError, "opacity"):
                export_standard_pptx([{"elements": [element]}])
            result = export_standard_pptx([{"elements": [element]}], unsupported="warn")
            self.assertEqual(len(Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes), 0)
            self.assertTrue(result.warnings)

    def test_absolute_and_relative_line_height(self):
        element = self.simple_text()
        element["font"] = {"size": 20, "line_height": 1.2}
        element["runs"] = [{"text": "第一行\n第二行"}]
        result = export_standard_pptx([{"elements": [element]}])
        paragraph = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes[0].text_frame.paragraphs[0]
        self.assertEqual(paragraph.line_spacing, 1.2)
        element["font"]["line_height"] = 24
        with self.assertRaisesRegex(NativeStandardExportError, "ambiguous"):
            export_standard_pptx([{"elements": [element]}])

    def test_aggregate_budgets_are_checked_before_building_office_objects(self):
        element = self.simple_text()
        element["runs"] = [{"text": "a" * 21}]
        with patch("services.native_standard_pptx.MAX_TEXT_CHARACTERS", 20), \
             patch("services.native_standard_pptx.Presentation", side_effect=AssertionError("Must preflight first")):
            with self.assertRaisesRegex(NativeStandardExportError, "characters limit"):
                export_standard_pptx([{"elements": [element]}])
        table = fixture()["slides"][1]["components"][1]["elements"][0]
        with patch("services.native_standard_pptx.MAX_TABLE_CELLS", 20):
            with self.assertRaisesRegex(NativeStandardExportError, "cells limit"):
                export_standard_pptx([{"elements": [table, table]}])
        cyclic = {}
        cyclic["elements"] = [cyclic]
        with self.assertRaisesRegex(NativeStandardExportError, "Cyclic"):
            export_standard_pptx([cyclic])

    def test_chart_labels_cannot_become_workbook_formulas_or_links(self):
        chart = fixture()["slides"][2]["components"][1]["elements"][0]
        chart["categories"] = ["=1+1", "+1+1", "@SUM(A1)", "https://example.invalid/label"]
        chart["series"] = [{"name": "=1+2", "values": [1, 2, 3, 4]}]
        result = export_standard_pptx([{"elements": [chart]}])
        with ZipFile(BytesIO(result.pptx_bytes)) as pptx:
            for name in pptx.namelist():
                if name.startswith("ppt/embeddings/"):
                    blob = pptx.read(name)
                    with ZipFile(BytesIO(blob)) as xlsx:
                        for part in xlsx.namelist():
                            if part.endswith(".xml"):
                                root = ET.fromstring(xlsx.read(part))
                                self.assertFalse(root.findall(".//s:f", NS), part)
                                self.assertFalse(root.findall(".//s:hyperlink", NS), part)
                            if part.endswith(".rels"):
                                root = ET.fromstring(xlsx.read(part))
                                self.assertFalse(any(rel.get("TargetMode") == "External" for rel in root), part)
                    rows = workbook_rows(blob)
                    self.assertEqual(rows[0], ["=1+2"])
                    self.assertEqual([row[0] for row in rows[1:]], chart["categories"])
        parsed = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes[0].chart
        self.assertEqual(parsed.series[0].name, "=1+2")
        self.assertEqual([category.label for category in parsed.plots[0].categories], chart["categories"])

    def test_malformed_inputs_raise_structured_errors_instead_of_internal_errors(self):
        cases = [
            {"type": []},
            {**self.simple_text(), "position": {"x": 10 ** 400, "y": 0}},
            {"type": "vector", "points": 3},
            {"type": "table", "columns": 3, "rows": [], "size": {"width": 100, "height": 100}},
            {"type": "chart", "series": 3, "chart_type": "bar", "categories": ["A"], "size": {"width": 100, "height": 100}},
            {"type": "container", "size": {"width": 100, "height": 100}, "border_radius": dict.fromkeys(("tl", "tr", "bl", "br"), [])},
        ]
        for element in cases:
            with self.subTest(element=element), self.assertRaises(NativeStandardExportError) as caught:
                export_standard_pptx([{"elements": [element]}])
            self.assertTrue(caught.exception.path.startswith("slides[0].ui.elements[0]"))

    def test_markdown_is_rejected_instead_of_exported_as_visible_delimiters(self):
        for value in ("**重点**", "_强调_", "A__重点__B"):
            element = self.simple_text()
            element["runs"] = [{"text": value}]
            with self.subTest(value=value), self.assertRaisesRegex(NativeStandardExportError, "Markdown"):
                export_standard_pptx([{"elements": [element]}])

    def test_chart_automatic_legend_and_color_fallback(self):
        element = fixture()["slides"][4]["components"][1]["elements"][0]
        element["legend"] = None
        element["colors"] = None
        element["color"] = "#102030"
        result = export_standard_pptx([{"elements": [element]}])
        chart = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes[0].chart
        self.assertTrue(chart.has_legend)
        self.assertEqual(str(chart.series[0].points[0].format.fill.fore_color.rgb), "102030")

    def test_manual_flow_and_container_positions_are_rejected(self):
        child = self.simple_text()
        child["__presenton_manual_position"] = True
        child["position"] = {"x": 120, "y": 70}
        for kind in ("container", "flex", "grid"):
            element = {"type": kind, "size": {"width": 400, "height": 200}, "columns": 1,
                       "child": child} if kind == "container" else {
                       "type": kind, "size": {"width": 400, "height": 200}, "columns": 1, "children": [child]}
            with self.subTest(kind=kind), self.assertRaisesRegex(NativeStandardExportError, "manual"):
                export_standard_pptx([{"elements": [element]}])

    def test_container_clipping_is_explicitly_rejected(self):
        child = self.simple_text()
        child["position"] = {"x": 80, "y": 0}
        element = {"type": "container", "size": {"width": 100, "height": 100},
                   "fill": {"color": "#FFFFFF"}, "child": child}
        with self.assertRaisesRegex(NativeStandardExportError, "Clipped"):
            export_standard_pptx([{"elements": [element]}])
        result = export_standard_pptx([{"elements": [element]}], unsupported="warn")
        self.assertEqual(len(Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes), 0)
        self.assertIn("Clipped", result.warnings[0].message)

    def test_polygon_closed_default_matches_current_canvas(self):
        for closed, should_close in ((None, True), (False, False), (True, True)):
            element = {"type": "vector", "closed": closed, "points": [
                       {"x": 0, "y": 0}, {"x": 100, "y": 0}, {"x": 50, "y": 100}]}
            result = export_standard_pptx([{"elements": [element]}])
            shape = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes[0]
            self.assertEqual(bool(shape._element.xpath(".//a:close")), should_close)

    def test_chart_strings_exceeding_workbook_limit_are_not_truncated(self):
        for field in ("category", "series"):
            chart = fixture()["slides"][2]["components"][1]["elements"][0]
            if field == "category":
                chart["categories"][0] = "文" * 32768
            else:
                chart["series"][0]["name"] = "文" * 32768
            with self.subTest(field=field), self.assertRaisesRegex(NativeStandardExportError, "32767"):
                export_standard_pptx([{"elements": [chart]}])

    def test_additional_malformed_styles_and_layouts_have_structured_errors(self):
        cases = [
            {"type": "flex", "direction": ["row"], "size": {"width": 400, "height": 200}, "children": [self.simple_text()]},
            {"type": "flex", "size": {"width": 400, "height": 200}, "children": [{**self.simple_text(), "type": ["text"]}]},
            {"type": "table", "size": {"width": 400, "height": 200}, "rows": [], "columns": [
             {"runs": [{"text": "x"}], "borders": {"top": {"width": "bad", "color": "FFFFFF"}}}]},
            {"type": "vector", "points": [{"x": 0, "y": 0}, {"x": 1, "y": 1}], "corner_radii": 3},
            {"type": "table", "size": {"width": 400, "height": 200}, "rows": [], "columns": [{"runs": 3}]},
        ]
        for element in cases:
            with self.subTest(element=element), self.assertRaises(NativeStandardExportError) as caught:
                export_standard_pptx([{"elements": [element]}])
            self.assertTrue(caught.exception.path.startswith("slides[0].ui.elements[0]"))

    def test_invalid_options_and_empty_deck_fail(self):
        for slides in ([], None, "no"):
            with self.subTest(slides=slides), self.assertRaises(NativeStandardExportError):
                export_standard_pptx(slides)
        with self.assertRaises(ValueError):
            export_standard_pptx([{}], unsupported="silent")

    @staticmethod
    def simple_text():
        return {"type": "text", "name": "test", "position": {"x": 0, "y": 0},
                "size": {"width": 100, "height": 40}, "runs": [{"text": "中文测试"}]}


if __name__ == "__main__":
    unittest.main()
