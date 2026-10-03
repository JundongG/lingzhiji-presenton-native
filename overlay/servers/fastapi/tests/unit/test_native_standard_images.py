# SPDX-License-Identifier: Apache-2.0
# Independent native-export add-on tests; all image pixels are synthetic.
"""Memory-only PNG/JPEG embedding, layout and security regressions."""
import copy
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation

from services.native_standard_pptx import EMU_PER_PIXEL, NativeStandardExportError, export_standard_pptx


def synthetic_image(fmt="PNG", size=(240, 120)):
    out = BytesIO()
    Image.new("RGB", size, (22, 166, 161)).save(out, format=fmt)
    return out.getvalue()


def image_element(**changes):
    return {"type": "image", "data": "synthetic:test", "position": {"x": 100, "y": 100},
            "size": {"width": 120, "height": 120}, "fit": "fill", **changes}


def image_references(value):
    if isinstance(value, dict):
        if value.get("type") == "image" and isinstance(value.get("data"), str):
            yield value["data"]
        for child in value.values():
            yield from image_references(child)
    elif isinstance(value, list):
        for child in value:
            yield from image_references(child)


class NativeStandardImageTests(unittest.TestCase):
    def export(self, element, blob=None):
        return export_standard_pptx([{"elements": [element]}], image_assets={element["data"]: blob or synthetic_image()})

    def test_png_and_jpeg_bytes_are_embedded_unchanged_as_native_pictures(self):
        for fmt in ("PNG", "JPEG"):
            blob = synthetic_image(fmt)
            result = self.export(image_element(), blob)
            shape = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes[0]
            self.assertEqual(shape.image.blob, blob)
            self.assertEqual((shape.left, shape.top, shape.width, shape.height), tuple(x * EMU_PER_PIXEL for x in (100, 100, 120, 120)))
            self.assertIn("p:pic", shape._element.xml)
            self.assertFalse(result.warnings)

    def test_cover_crop_and_focus_use_percent_semantics(self):
        for focus, left, right in ((0, 0, .5), (50, .25, .25), (100, .5, 0)):
            result = self.export(image_element(fit="cover", focus_x=focus))
            shape = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes[0]
            self.assertAlmostEqual(shape.crop_left, left)
            self.assertAlmostEqual(shape.crop_right, right)
            self.assertEqual(shape.crop_top, 0)
            self.assertEqual(shape.crop_bottom, 0)

    def test_default_contain_preserves_aspect_and_native_geometry(self):
        result = self.export(image_element(fit=None))
        shape = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes[0]
        self.assertEqual((shape.left, shape.top, shape.width, shape.height), tuple(x * EMU_PER_PIXEL for x in (100, 130, 120, 60)))

    def test_native_rounding_alpha_and_flips_do_not_rewrite_image_bytes(self):
        blob = synthetic_image()
        result = self.export(image_element(border_radius={key: 12 for key in ("tl", "tr", "bl", "br")},
                                           opacity=.25, flip_h=True), blob)
        shape = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes[0]
        self.assertEqual(shape.image.blob, blob)
        self.assertIn('prst="roundRect"', shape._element.xml)
        self.assertIn('amt="25000"', shape._element.xml)
        self.assertIn('flipH="1"', shape._element.xml)

    def test_missing_reference_never_loads_a_file_or_url(self):
        with patch("socket.socket", side_effect=AssertionError("No network")), \
             patch.object(Path, "read_bytes", side_effect=AssertionError("No asset reads")):
            for reference in ("https://example.invalid/private.png", "/etc/passwd", "file:///etc/passwd"):
                with self.subTest(reference=reference), self.assertRaises(NativeStandardExportError) as caught:
                    export_standard_pptx([{"elements": [image_element(data=reference)]}])
                self.assertEqual(caught.exception.path, "slides[0].ui.elements[0].data")
                self.assertEqual(caught.exception.code, "missing_image_asset")

    def test_assets_are_bytes_not_paths_streams_or_callbacks(self):
        for value in ("/tmp/fake.png", BytesIO(b"x"), lambda: b"x", bytearray(b"x")):
            with self.subTest(value=type(value).__name__), self.assertRaises(NativeStandardExportError):
                export_standard_pptx([{}], image_assets={"x": value})

    def test_invalid_unsupported_animated_and_rotated_exif_images_are_rejected(self):
        animation = BytesIO()
        Image.new("RGB", (5, 5), "red").save(animation, format="PNG", save_all=True,
                                             append_images=[Image.new("RGB", (5, 5), "blue")])
        rotated = BytesIO()
        exif = Image.Exif(); exif[274] = 6
        Image.new("RGB", (5, 5)).save(rotated, format="JPEG", exif=exif)
        for blob in (b"not an image", b'<svg xmlns="http://www.w3.org/2000/svg"/>', synthetic_image("BMP"), animation.getvalue(), rotated.getvalue()):
            with self.subTest(size=len(blob)), self.assertRaises(NativeStandardExportError):
                self.export(image_element(), blob)

    def test_byte_count_pixel_and_side_budgets_are_enforced(self):
        blob = synthetic_image()
        cases = (("MAX_IMAGE_BYTES", len(blob)-1, {"x": blob}),
                 ("MAX_TOTAL_IMAGE_BYTES", len(blob)*2-1, {"x": blob, "y": blob}),
                 ("MAX_IMAGE_ASSETS", 1, {"x": blob, "y": blob}),
                 ("MAX_IMAGE_PIXELS", 100, {"x": blob}),
                 ("MAX_TOTAL_IMAGE_PIXELS", 50000, {"x": blob, "y": blob}),
                 ("MAX_IMAGE_SIDE", 100, {"x": blob}))
        for constant, limit, assets in cases:
            with self.subTest(constant=constant), patch("services.native_standard_pptx." + constant, limit), self.assertRaises(NativeStandardExportError) as caught:
                export_standard_pptx([{}], image_assets=assets)
            self.assertEqual(caught.exception.code, "export_limit")

    def test_unsupported_picture_styles_are_explicit(self):
        for options in ({"clip_path": "M0 0"}, {"clippath": "M0 0"}, {"clipPath": "M0 0"}, {"crop_scale": 2}, {"color": "#FF0000"},
                        {"focus_x": 101}, {"border_radius": {"tl": 1, "tr": 2, "bl": 1, "br": 2}},
                        {"fit": "contain", "border_radius": 10}, {"fit": "contain", "rotation": 10, "focus_x": 0}):
            with self.subTest(options=options), self.assertRaises(NativeStandardExportError):
                self.export(image_element(**options))

    def test_embedding_is_in_memory_and_does_not_mutate_inputs(self):
        element = image_element(fit="cover")
        original = copy.deepcopy(element)
        blob = synthetic_image()
        with patch("socket.socket", side_effect=AssertionError("No network")), \
             patch.object(tempfile, "mkstemp", side_effect=AssertionError("No temp files")):
            self.export(element, blob)
        self.assertEqual(element, original)

    def test_rounded_rectangular_vectors_and_group_flow_are_native(self):
        vector = {"type": "vector", "points": [{"x": 0, "y": 0}, {"x": 80, "y": 0}, {"x": 80, "y": 40}, {"x": 0, "y": 40}],
                  "corner_radii": [8]*4, "fill": {"color": "#123456"}}
        group = {"type": "group", "size": {"width": 80, "height": 40}, "children": [vector]}
        ui = {"elements": [{"type": "grid", "position": {"x": 100, "y": 100}, "size": {"width": 180, "height": 40},
                            "columns": 2, "gap": 20, "children": [group, group]}]}
        result = export_standard_pptx([ui])
        shapes = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes
        self.assertEqual([shape.left for shape in shapes], [100*EMU_PER_PIXEL, 200*EMU_PER_PIXEL])
        self.assertTrue(all('prst="roundRect"' in shape._element.xml for shape in shapes))

    def test_contain_flips_reflect_picture_placement_in_the_original_frame(self):
        result = self.export(image_element(fit="contain", focus_y=0, flip_v=True))
        shape = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes[0]
        self.assertEqual(shape.top, 160*EMU_PER_PIXEL)
        self.assertEqual(shape.height, 60*EMU_PER_PIXEL)
        self.assertIn('flipV="1"', shape._element.xml)
        result = self.export(image_element(fit="contain", focus_x=0, flip_h=True), synthetic_image(size=(120, 240)))
        shape = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes[0]
        self.assertEqual(shape.left, 160*EMU_PER_PIXEL)
        self.assertEqual(shape.width, 60*EMU_PER_PIXEL)

    def test_vector_center_rotation_and_mirroring_are_native(self):
        element = {"type": "vector", "points": [{"x": 100, "y": 100}, {"x": 180, "y": 100},
                   {"x": 180, "y": 140}, {"x": 100, "y": 140}], "rotation": 30, "flip_h": True,
                   "corner_radii": [8]*4}
        result = export_standard_pptx([{"elements": [element]}])
        shape = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes[0]
        self.assertEqual(shape.rotation, 30)
        self.assertEqual((shape.left, shape.top, shape.width, shape.height), tuple(v*EMU_PER_PIXEL for v in (100, 100, 80, 40)))
        self.assertIn('flipH="1"', shape._element.xml)

    def test_container_bounded_optional_text_size_is_resolved(self):
        ui = {"elements": [{"type": "container", "position": {"x": 100, "y": 100}, "size": {"width": 48, "height": 48},
                            "alignment": {"horizontal": "center", "vertical": "middle"}, "child": {
                            "type": "text", "font": {"size": 16, "line_height": 1.2}, "runs": [{"text": "01"}],
                            "alignment": {"horizontal": "center", "vertical": "middle"}}}]}
        result = export_standard_pptx([ui])
        shape = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes[0]
        self.assertEqual(shape.width, 48*EMU_PER_PIXEL)
        self.assertEqual(shape.height, round(19.2*EMU_PER_PIXEL))
        self.assertEqual(shape.text, "01")

    def test_unsized_lists_are_explicitly_rejected_and_text_matches_canvas_estimator(self):
        for items in ([[{"text": str(index)}] for index in range(5)], 3, [3]):
            ui = {"elements": [{"type": "container", "size": {"width": 100, "height": 250},
                                "child": {"type": "text-list", "items": items}}]}
            with self.subTest(items=items), self.assertRaises(NativeStandardExportError):
                export_standard_pptx([ui])
        ui = {"elements": [{"type": "container", "size": {"width": 100, "height": 250},
                            "child": {"type": "text", "runs": [{"text": "a" * 111}]}}]}
        result = export_standard_pptx([ui])
        shape = Presentation(BytesIO(result.pptx_bytes)).slides[0].shapes[0]
        self.assertEqual(shape.height, round(11 * 18 * 1.15 * EMU_PER_PIXEL))

    def test_all_eleven_standard_structural_prototypes_with_synthetic_assets(self):
        path = Path(__file__).parents[4] / "templates" / "standard" / "template.json"
        template = json.loads(path.read_text(encoding="utf-8"))
        self.assertEqual(len(template["layouts"]), 11)
        blob = synthetic_image()
        for ui in template["layouts"]:
            with self.subTest(layout=ui["id"]):
                # Structural regression only: each original reference receives
                # a synthetic PNG. This does not validate actual asset loading,
                # authorize SVG conversion, or claim a generated-deck UI test.
                refs = set(image_references(ui))
                result = export_standard_pptx([ui], image_assets={ref: blob for ref in refs})
                self.assertEqual(len(Presentation(BytesIO(result.pptx_bytes)).slides), 1)
                self.assertFalse(any(warning.code == "missing_image_asset" for warning in result.warnings))


if __name__ == "__main__":
    unittest.main()
